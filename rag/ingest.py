"""Document processing pipeline: extraction -> cleaning -> chunking."""
from __future__ import annotations

import csv
import hashlib
import io
import re
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".csv", ".docx"}


@dataclass
class Document:
    doc_id: str
    source: str
    text: str
    metadata: dict = field(default_factory=dict)


@dataclass
class Chunk:
    chunk_id: str
    doc_id: str
    source: str
    text: str
    position: int
    metadata: dict = field(default_factory=dict)


# ---------------------------------------------------------------- extraction
def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        pages.append(f"[Page {i}]\n{page.extract_text() or ''}")
    return "\n\n".join(pages)


def _extract_docx(data: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - optional dependency
        raise RuntimeError("Install python-docx to ingest .docx files") from exc
    document = docx.Document(io.BytesIO(data))
    return "\n".join(p.text for p in document.paragraphs)


def _extract_csv(data: bytes) -> str:
    """Render a CSV as readable text so its content is searchable."""
    text = data.decode("utf-8", errors="ignore")
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        return ""
    header, body = rows[0], rows[1:]
    lines = [f"Table with columns: {', '.join(header)}"]
    for row in body:
        lines.append("; ".join(f"{h}: {v}" for h, v in zip(header, row)))
    return "\n".join(lines)


def extract_text(name: str, data: bytes) -> str:
    ext = Path(name).suffix.lower()
    if ext == ".pdf":
        return _extract_pdf(data)
    if ext == ".docx":
        return _extract_docx(data)
    if ext == ".csv":
        return _extract_csv(data)
    if ext in {".txt", ".md"}:
        return data.decode("utf-8", errors="ignore")
    raise ValueError(f"Unsupported file type: {ext}")


# ------------------------------------------------------------------ cleaning
def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", " ", text)  # control chars
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)  # de-hyphenate line breaks
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# ------------------------------------------------------------------ chunking
def _split_sentences(paragraph: str) -> list[str]:
    parts = re.split(r"(?<=[.!?؟])\s+", paragraph)
    return [p.strip() for p in parts if p.strip()]


def chunk_text(text: str, chunk_size: int = 800, overlap: int = 150) -> list[str]:
    """Paragraph/sentence-aware chunking with character overlap.

    Paragraphs are packed together until ``chunk_size`` characters; long
    paragraphs are split on sentence boundaries. Each new chunk starts with the
    tail of the previous one so facts spanning a boundary are not lost.
    """
    units: list[str] = []
    for para in text.split("\n\n"):
        para = para.strip()
        if not para:
            continue
        if len(para) <= chunk_size:
            units.append(para)
        else:
            units.extend(_split_sentences(para))

    chunks: list[str] = []
    current = ""
    for unit in units:
        if current and len(current) + len(unit) + 2 > chunk_size:
            chunks.append(current.strip())
            tail = current[-overlap:] if overlap else ""
            # start the overlap at a word boundary
            tail = tail[tail.find(" ") + 1 :] if " " in tail else tail
            current = f"{tail} {unit}" if tail else unit
        else:
            current = f"{current}\n\n{unit}" if current else unit
        while len(current) > chunk_size * 1.5:  # a single huge unit
            chunks.append(current[:chunk_size].strip())
            current = current[chunk_size - overlap :]
    if current.strip():
        chunks.append(current.strip())
    return chunks


# ------------------------------------------------------------------ pipeline
def load_document(name: str, data: bytes) -> Document:
    text = clean_text(extract_text(name, data))
    doc_id = hashlib.md5(f"{name}:{len(data)}".encode() + data[:2048]).hexdigest()[:12]
    title = Path(name).stem.replace("_", " ").replace("-", " ").title()
    return Document(doc_id=doc_id, source=Path(name).name, text=text, metadata={"title": title})


def load_directory(directory: str | Path) -> list[Document]:
    docs = []
    for path in sorted(Path(directory).iterdir()):
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
            docs.append(load_document(path.name, path.read_bytes()))
    return docs


def chunk_documents(docs: list[Document], chunk_size: int = 800, overlap: int = 150) -> list[Chunk]:
    chunks = []
    for doc in docs:
        for i, piece in enumerate(chunk_text(doc.text, chunk_size, overlap)):
            chunks.append(
                Chunk(
                    chunk_id=f"{doc.doc_id}-{i}",
                    doc_id=doc.doc_id,
                    source=doc.source,
                    text=piece,
                    position=i,
                    metadata=dict(doc.metadata),
                )
            )
    return chunks
